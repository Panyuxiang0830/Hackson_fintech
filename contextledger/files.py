"""Extract supported native documents without dropping their source files."""

from __future__ import annotations

import csv
from pathlib import Path


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return "\n\n".join(page.extract_text() or "" for page in pdf.pages).strip()
    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        parts.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(parts).strip()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            parts = []
            for sheet in workbook:
                parts.append(f"Sheet: {sheet.title}")
                parts.extend(" | ".join("" if value is None else str(value) for value in row) for row in sheet.values)
            return "\n".join(parts).strip()
        finally:
            workbook.close()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle)).strip()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8").strip()
    raise ValueError(f"unsupported document type: {path.name}")
